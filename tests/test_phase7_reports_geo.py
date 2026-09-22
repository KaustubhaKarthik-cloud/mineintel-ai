"""Phase 7 extensions — geological/combined reports, PDF/Excel, rejected exclusion."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.models import Document, FactStatus, GeologicalFact, IndexStatus
from app.reports.export_excel import write_report_excel
from app.reports.export_pdf import write_report_pdf
from app.reports.geo_sections import collect_geological_payload
from app.reports.service import build_full_report, export_report_excel, export_report_pdf


def _doc(db, name="phase7_geo.pdf") -> Document:
    d = Document(
        filename=name,
        original_filename=name,
        file_path=f"/tmp/{name}",
        file_type="pdf",
        status="extracted",
        index_status=IndexStatus.INDEXED.value,
        document_category="Geological Exploration Report",
        version=1,
        meta={
            "g1_classification": {
                "domain": "geological_exploration",
                "label": "Geological Exploration Report",
            }
        },
    )
    db.add(d)
    db.flush()
    return d


def _geo_fact(db, doc: Document, **kwargs) -> GeologicalFact:
    defaults = {
        "document_id": doc.id,
        "domain": "geological",
        "status": FactStatus.REVIEW_REQUIRED.value,
        "extraction_confidence": 0.7,
        "source_page": 10,
        "evidence_text": "structured geological evidence",
    }
    defaults.update(kwargs)
    row = GeologicalFact(**defaults)
    db.add(row)
    db.flush()
    return row


def test_empty_geological_report(db_session, tmp_path, monkeypatch):
    from app.utils import files as files_mod

    monkeypatch.setattr(files_mod, "documents_root", lambda: tmp_path)
    (tmp_path / "reports").mkdir(parents=True)
    report = build_full_report(
        db_session,
        title="Empty geo",
        domain="geological",
        generated_by="tester",
    )
    assert report.status == "ready"
    assert "Insufficient" in (report.content or "") or "No geological" in (report.content or "")
    assert "Disclaimer" in (report.content or "") or "disclaimer" in (report.content or "").lower()


def test_geological_report_includes_formations_and_marks_review(db_session, tmp_path, monkeypatch):
    from app.utils import files as files_mod

    monkeypatch.setattr(files_mod, "documents_root", lambda: tmp_path)
    (tmp_path / "reports").mkdir(parents=True)
    doc = _doc(db_session)
    _geo_fact(
        db_session,
        doc,
        metric_kind="formation",
        geological_formation="Raniganj Formation",
        evidence_text="The Raniganj Formation hosts coal seams.",
        original_value="Raniganj Formation",
        source_page=11,
    )
    _geo_fact(
        db_session,
        doc,
        metric_kind="resource_quantity",
        seam_name="R4",
        seam_status="named",
        original_value="1.67",
        original_unit="MT",
        source_page=31,
        evidence_text="Seam R4 contributes 1.67 MT of inferred coal resources.",
    )
    _geo_fact(
        db_session,
        doc,
        metric_kind="resource_quantity",
        seam_name="R4",
        original_value="99.9",
        original_unit="MT",
        status=FactStatus.REJECTED.value,
        source_page=99,
        evidence_text="Rejected 99.9 MT.",
    )
    report = build_full_report(
        db_session,
        domain="geological",
        document_ids=[doc.id],
        title="Geo report",
    )
    html = report.content or ""
    assert "Raniganj Formation" in html
    assert "1.67" in html
    assert "99.9" not in html
    assert "REVIEW REQUIRED" in html or "review" in html.lower()
    assert (report.parameters or {}).get("rejected_excluded") >= 1
    assert (report.parameters or {}).get("pending_verification") is True
    # Provenance has page
    pages = [p.get("page") for p in (report.provenance or []) if p.get("page") is not None]
    assert 11 in pages or 31 in pages


def test_combined_report_domain(db_session, tmp_path, monkeypatch):
    from app.utils import files as files_mod

    monkeypatch.setattr(files_mod, "documents_root", lambda: tmp_path)
    (tmp_path / "reports").mkdir(parents=True)
    doc = _doc(db_session, "combo.pdf")
    _geo_fact(
        db_session,
        doc,
        metric_kind="formation",
        geological_formation="Barakar Formation",
        evidence_text="Barakar Formation is present.",
        original_value="Barakar Formation",
    )
    report = build_full_report(db_session, domain="combined", document_ids=[doc.id])
    assert (report.parameters or {}).get("domain") == "combined"
    assert "Barakar Formation" in (report.content or "")
    assert "Disclaimer" in (report.content or "") or "disclaimer" in (report.content or "").lower()


def test_pdf_export_creates_file(tmp_path):
    out = tmp_path / "r.pdf"
    write_report_pdf(
        title="Test Report",
        html_content="<h1>Hello</h1><p>Value 1.67 MT on page 31</p>",
        output_path=out,
        meta={"generated_at": "2026-01-01", "domain": "geological", "warnings": ["pending"]},
    )
    assert out.is_file()
    assert out.stat().st_size > 200
    import fitz

    doc = fitz.open(str(out))
    assert len(doc) >= 1
    text = "".join(page.get_text() for page in doc)
    doc.close()
    assert "MineIntel" in text
    assert "1.67" in text
    assert "Page 1" in text


def test_excel_export_sheets(tmp_path):
    out = tmp_path / "r.xlsx"
    write_report_excel(
        output_path=out,
        summary={
            "title": "T",
            "domain": "geological",
            "generated_at": "now",
            "source_documents": ["a.pdf"],
            "review_required_count": 2,
            "rejected_excluded": 1,
            "disclaimer": "test disclaimer",
        },
        geological_facts=[
            {
                "id": "f1",
                "metric_kind": "resource_quantity",
                "seam_label": "Seam R4",
                "display_value": "1.67",
                "display_unit": "MT",
                "status": "review_required",
                "document_name": "a.pdf",
                "source_page": 31,
                "requires_human_verification": True,
            }
        ],
        resources=[{"seam": "Seam R4", "value": "1.67", "unit": "MT", "page": 31, "status": "review_required"}],
        evidence=[
            {
                "fact_id": "f1",
                "document_name": "a.pdf",
                "page": 31,
                "status_label": "REVIEW REQUIRED",
                "value": "1.67",
                "evidence_text": "1.67 MT",
            }
        ],
    )
    assert out.is_file()
    wb = load_workbook(out)
    names = set(wb.sheetnames)
    assert "Summary" in names
    assert "Geological_Facts" in names
    assert "Evidence_Sources" in names
    assert "Resources" in names
    assert "Audit_Review_Status" in names
    flat = []
    for row in wb["Geological_Facts"].iter_rows(values_only=True):
        flat.extend(row)
    assert "1.67" in flat
    assert "review_required" in flat


def test_collect_payload_excludes_rejected(db_session):
    doc = _doc(db_session, "rej.pdf")
    _geo_fact(
        db_session,
        doc,
        metric_kind="formation",
        geological_formation="Raniganj Formation",
        evidence_text="Raniganj Formation present.",
    )
    _geo_fact(
        db_session,
        doc,
        metric_kind="formation",
        geological_formation="Fake Formation",
        status=FactStatus.REJECTED.value,
        evidence_text="Fake Formation rejected.",
    )
    payload = collect_geological_payload(db_session, document_ids=[doc.id])
    names = {f["formation_name"] for f in payload["formations"]}
    assert "Raniganj Formation" in names
    assert "Fake Formation" not in names
    assert payload["rejected_excluded"] >= 1


def test_api_geological_pdf_excel(tmp_path, monkeypatch):
    """API smoke with auth-required app (mirrors phase7_final fixture)."""
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from app.auth.security import create_access_token, hash_password
    from app.database import Base, get_db
    from app.main import create_app
    from app.models import User, UserRole, UserStatus
    from app.utils import files as files_mod

    monkeypatch.setenv("USE_SQLITE", "true")
    monkeypatch.setenv("AUTH_REQUIRED", "true")
    monkeypatch.setenv("AUTH_SEED_USERS", "false")
    monkeypatch.setenv("ASSISTANT_LLM_PROVIDER", "mock")
    monkeypatch.setenv("DOCUMENTS_DIR", str(tmp_path / "documents"))
    (tmp_path / "documents" / "original").mkdir(parents=True)
    (tmp_path / "documents" / "reports").mkdir(parents=True)
    monkeypatch.setattr(files_mod, "documents_root", lambda: tmp_path / "documents")

    from app.config import get_settings

    get_settings.cache_clear()

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSession = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def _override():
        db = TestingSession()
        try:
            yield db
        finally:
            db.close()

    app = create_app()
    app.dependency_overrides[get_db] = _override
    client = TestClient(app)

    db = TestingSession()
    analyst = User(
        username="analyst",
        display_name="Analyst",
        password_hash=hash_password("analyst123"),
        role=UserRole.ANALYST.value,
        status=UserStatus.ACTIVE.value,
    )
    reviewer = User(
        username="reviewer",
        display_name="Reviewer",
        password_hash=hash_password("reviewer123"),
        role=UserRole.REVIEWER.value,
        status=UserStatus.ACTIVE.value,
    )
    db.add_all([analyst, reviewer])
    db.flush()
    doc = Document(
        filename="g.pdf",
        original_filename="g.pdf",
        file_path=str(tmp_path / "g.pdf"),
        file_type="pdf",
        status="extracted",
        version=1,
    )
    db.add(doc)
    db.flush()
    db.add(
        GeologicalFact(
            document_id=doc.id,
            metric_kind="resource_quantity",
            seam_name="R4",
            seam_status="named",
            original_value="1.67",
            original_unit="MT",
            status=FactStatus.REVIEW_REQUIRED.value,
            source_page=31,
            evidence_text="Seam R4 contributes 1.67 MT of inferred coal resources.",
            extraction_confidence=0.8,
        )
    )
    db.commit()
    analyst_id, reviewer_id, doc_id = analyst.id, reviewer.id, doc.id
    db.close()

    def _auth(uid: str, role: str):
        return {
            "Authorization": f"Bearer {create_access_token(user_id=uid, username=role, role=role)}"
        }

    try:
        gen = client.post(
            "/api/reports/generate",
            headers=_auth(analyst_id, "analyst"),
            json={"domain": "geological", "document_ids": [doc_id], "title": "Geo API"},
        )
        assert gen.status_code == 200, gen.text
        rid = gen.json()["id"]
        assert gen.json().get("domain") == "geological" or gen.json()["parameters"].get("domain") == "geological"

        pdf = client.get(f"/api/reports/{rid}/pdf", headers=_auth(analyst_id, "analyst"))
        assert pdf.status_code == 200
        assert pdf.content[:4] == b"%PDF"

        xlsx = client.get(f"/api/reports/{rid}/excel", headers=_auth(analyst_id, "analyst"))
        assert xlsx.status_code == 200
        assert xlsx.content[:2] == b"PK"

        # USER (legacy analyst) cannot access admin-only validation
        denied = client.get(
            "/api/validation/conflicts",
            headers=_auth(analyst_id, "analyst"),
        )
        assert denied.status_code == 403
        ok_read = client.get(f"/api/reports/{rid}/pdf", headers=_auth(reviewer_id, "reviewer"))
        assert ok_read.status_code == 200
        # Reviewer (→ admin) can generate
        ok_gen = client.post(
            "/api/reports/generate",
            headers=_auth(reviewer_id, "reviewer"),
            json={"domain": "geological", "document_ids": [doc_id], "title": "Geo Admin"},
        )
        assert ok_gen.status_code == 200
    finally:
        app.dependency_overrides.clear()
        get_settings.cache_clear()


def test_export_helpers_from_report_row(db_session, tmp_path, monkeypatch):
    from app.utils import files as files_mod

    monkeypatch.setattr(files_mod, "documents_root", lambda: tmp_path)
    (tmp_path / "reports").mkdir(parents=True)
    report = build_full_report(db_session, domain="mining", title="Mining empty")
    pdf = export_report_pdf(db_session, report)
    xlsx = export_report_excel(db_session, report)
    assert Path(pdf).is_file()
    assert Path(xlsx).is_file()
