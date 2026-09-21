"""Excel export for MineIntel reports using openpyxl."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font


def _ws_write(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(["" if v is None else v for v in row])


def write_report_excel(
    *,
    output_path: Path,
    summary: dict[str, Any],
    mining_facts: Optional[list[dict[str, Any]]] = None,
    geological_facts: Optional[list[dict[str, Any]]] = None,
    conflicts: Optional[list[dict[str, Any]]] = None,
    evidence: Optional[list[dict[str, Any]]] = None,
    resources: Optional[list[dict[str, Any]]] = None,
) -> Path:
    """Write only sheets that have relevant data."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Summary (always)
    ws = wb.active
    ws.title = "Summary"
    _ws_write(
        ws,
        ["Field", "Value"],
        [
            ["title", summary.get("title")],
            ["domain", summary.get("domain")],
            ["report_type", summary.get("report_type")],
            ["generated_at", summary.get("generated_at")],
            ["generated_by", summary.get("generated_by")],
            ["documents", ", ".join(summary.get("source_documents") or [])],
            ["fact_count", summary.get("fact_count")],
            ["geo_fact_count", summary.get("geo_fact_count")],
            ["open_conflicts", summary.get("open_conflicts")],
            ["review_required_count", summary.get("review_required_count")],
            ["rejected_excluded", summary.get("rejected_excluded")],
            ["pending_verification", summary.get("pending_verification")],
            ["disclaimer", summary.get("disclaimer")],
        ],
    )
    if summary.get("warnings"):
        ws.append([])
        ws.append(["Warnings"])
        for w in summary["warnings"]:
            ws.append([w])

    if mining_facts:
        ws_f = wb.create_sheet("Facts_Metrics")
        _ws_write(
            ws_f,
            [
                "entity",
                "metric",
                "value",
                "unit",
                "period",
                "document",
                "document_id",
                "page",
                "fact_id",
                "status",
            ],
            [
                [
                    f.get("entity"),
                    f.get("metric"),
                    f.get("value"),
                    f.get("unit"),
                    f.get("period") or f.get("fiscal_year"),
                    f.get("document_name"),
                    f.get("document_id"),
                    f.get("page"),
                    f.get("fact_id") or f.get("chunk_id"),
                    f.get("status") or "verified_structured",
                ]
                for f in mining_facts
            ],
        )

    if conflicts:
        ws_c = wb.create_sheet("Validation_Conflicts")
        _ws_write(
            ws_c,
            ["entity", "field", "period", "severity", "status", "description"],
            [
                [
                    c.get("entity_name") or c.get("entity"),
                    c.get("field_name") or c.get("field"),
                    c.get("period"),
                    c.get("severity"),
                    c.get("status"),
                    c.get("description"),
                ]
                for c in conflicts
            ],
        )

    if geological_facts:
        ws_g = wb.create_sheet("Geological_Facts")
        _ws_write(
            ws_g,
            [
                "fact_id",
                "metric_kind",
                "seam",
                "formation",
                "borehole",
                "value",
                "unit",
                "status",
                "document",
                "document_id",
                "page",
                "requires_review",
            ],
            [
                [
                    g.get("id") or g.get("fact_id"),
                    g.get("metric_kind"),
                    g.get("seam_label") or g.get("seam_name"),
                    g.get("formation_name") or g.get("geological_formation"),
                    g.get("borehole_id"),
                    g.get("display_value") or g.get("original_value") or g.get("value"),
                    g.get("display_unit") or g.get("original_unit") or g.get("unit"),
                    g.get("status"),
                    g.get("document_name"),
                    g.get("document_id") or g.get("source_document_id"),
                    g.get("source_page") or g.get("page"),
                    g.get("requires_human_verification"),
                ]
                for g in geological_facts
            ],
        )

    if resources:
        ws_r = wb.create_sheet("Resources")
        _ws_write(
            ws_r,
            ["seam", "value", "unit", "page", "status", "document", "fact_id"],
            [
                [
                    r.get("seam"),
                    r.get("value"),
                    r.get("unit"),
                    r.get("page"),
                    r.get("status"),
                    r.get("document_name"),
                    r.get("fact_id"),
                ]
                for r in resources
            ],
        )

    if evidence:
        ws_e = wb.create_sheet("Evidence_Sources")
        _ws_write(
            ws_e,
            [
                "fact_id",
                "document",
                "document_id",
                "page",
                "status",
                "metric",
                "value",
                "evidence",
            ],
            [
                [
                    e.get("fact_id"),
                    e.get("document_name"),
                    e.get("document_id"),
                    e.get("page"),
                    e.get("status_label") or e.get("status"),
                    e.get("metric_kind") or e.get("metric"),
                    e.get("value"),
                    e.get("evidence_text") or e.get("evidence"),
                ]
                for e in evidence
            ],
        )

    # Audit / review status sheet
    ws_a = wb.create_sheet("Audit_Review_Status")
    _ws_write(
        ws_a,
        ["metric", "count"],
        [
            ["review_required_count", summary.get("review_required_count") or 0],
            ["rejected_excluded", summary.get("rejected_excluded") or 0],
            ["pending_verification", summary.get("pending_verification")],
            ["open_conflicts", summary.get("open_conflicts") or 0],
            ["verified_note", "Rejected facts are never listed as verified"],
        ],
    )

    wb.save(str(output_path))
    return output_path
